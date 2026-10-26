"""
generate_pptx_2026_10_27.py — 幼儿园大班今日新闻 2026-10-27 PPTX 生成器
3条新闻，每条一页，无封面，无天气页。
包含简单的动画效果。
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

# ---------- Palette ----------
PALETTE = {
    "bg":       RGBColor(0xFF, 0xFA, 0xF0),
    "ink":      RGBColor(0x25, 0x45, 0x37),
    "accent":   RGBColor(0xBC, 0x4F, 0x32),
    "chip_bg":  RGBColor(0xE8, 0xF0, 0xE5),
    "card_bg":  RGBColor(0xFF, 0xFF, 0xFF),
    "muted":    RGBColor(0x52, 0x66, 0x5B),
    "white":    RGBColor(0xFF, 0xFF, 0xFF),
}

FONT_HEAD = "Microsoft YaHei"
FONT_BODY = "Microsoft YaHei"

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ASSETS_DIR = os.path.join(SCRIPT_DIR, "daily-news", "2026-10-27", "assets")

NEWS = [
    {
        "chip": "1 / 3",
        "topic": "太空里的互联网",
        "title": "第28组卫星从海南飞上天",
        "body1": "10月25日，长征八号甲火箭从海南起飞，把第28组卫星送进太空。",
        "body2": "这些卫星在天上组成一张大网，让偏远地方也能上网。",
        "purpose": "以后在很远很远的地方，也能用卫星上网课、看视频。",
        "image": os.path.join(ASSETS_DIR, "satellite-launch.png"),
    },
    {
        "chip": "2 / 3",
        "topic": "体操世锦赛",
        "title": "中国体操队时隔20年再夺世界冠军",
        "body1": "在荷兰鹿特丹，中国女子体操队拿了世锦赛团体冠军。",
        "body2": "这是中国队时隔20年再夺这个冠军，上次夺冠还是2006年。",
        "purpose": "团队合作和坚持练习，就能离梦想更近。",
        "image": os.path.join(ASSETS_DIR, "gymnastics-team-1.png"),
    },
    {
        "chip": "3 / 3",
        "topic": "世界雪豹日",
        "title": "雪豹的大尾巴像围巾",
        "body1": "住在高山上",
        "body2": "厚厚的皮毛",
        "purpose": "大大的尾巴：走路稳，还能保暖",
        "image": os.path.join(ASSETS_DIR, "snow-leopard-photo.jpg"),
        "animation": os.path.join(ASSETS_DIR, "snow-leopard.gif"),
        "special": "snow_leopard",
    },
]


def add_text(slide, text, x, y, w, h, *,
             size=14, bold=False, color=None, align=PP_ALIGN.LEFT,
             font=None, font_ea=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font or FONT_BODY
    r.font.size = Pt(size)
    r.font.bold = bold
    if color is not None:
        r.font.color.rgb = color
    rPr = r._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {"typeface": font_ea or FONT_BODY})
        rPr.append(ea)
    else:
        ea.set("typeface", font_ea or FONT_BODY)
    return tb


def add_rect(slide, x, y, w, h, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
    shp.shadow.inherit = False
    return shp


def add_rounded_rect(slide, x, y, w, h, fill):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_animation(slide, shape, delay_ms=0, duration_ms=500):
    """Add a simple fade-in animation to a shape."""
    timing = slide._element.find(qn('p:timing'))
    if timing is None:
        timing = parse_xml(
            '<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
            '<p:tnLst><p:par><p:cT id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
            '<p:childTnLst><p:seq concurrent="1" nextAc="seek">'
            '<p:cT id="2" dur="indefinite" nodeType="mainSeq">'
            '<p:childTnLst></p:childTnLst></p:cT></p:seq></p:childTnLst></p:cT></p:par></p:tnLst></p:timing>'
        )
        slide._element.append(timing)

    # Get the main sequence child list
    seq = timing.find(qn('p:tnLst')).find(qn('p:par')).find(qn('p:cT')).find(qn('p:childTnLst')).find(qn('p:seq'))
    child_tn_lst = seq.find(qn('p:cT')).find(qn('p:childTnLst'))

    # Create animation for this shape
    shape_id = shape.shape_id
    anim_xml = f'''<p:par xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
      <p:cT id="{3 + len(child_tn_lst)}" fill="hold">
        <p:stCondLst><p:cond delay="{delay_ms}"/></p:stCondLst>
        <p:childTnLst>
          <p:set>
            <p:cBhvr><p:cT id="{4 + len(child_tn_lst)}" dur="1" fill="hold">
              <p:stCondLst><p:cond delay="0"/></p:stCondLst>
            </p:cT><p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl>
            <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
          </p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>
          <p:animEffect transition="in" filter="fade">
            <p:cBhvr><p:cT id="{5 + len(child_tn_lst)}" dur="{duration_ms}"/>
            <p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl></p:cBhvr>
          </p:animEffect>
        </p:childTnLst>
      </p:cT>
    </p:par>'''

    child_tn_lst.append(parse_xml(anim_xml))


def build_snow_leopard_slide(prs, news, slide_num):
    """少字、图片为主的雪豹页；GIF在PowerPoint放映模式中自动播放。"""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = add_rect(s, 0, 0, 13.333, 7.5, PALETTE["bg"])

    chip = add_rounded_rect(s, 0.6, 0.42, 1.1, 0.45, PALETTE["chip_bg"])
    add_text(s, news["chip"], 0.6, 0.42, 1.1, 0.45,
             size=14, bold=True, color=PALETTE["ink"], align=PP_ALIGN.CENTER)
    topic = add_text(s, news["topic"], 1.9, 0.42, 3.0, 0.45,
                     size=14, color=PALETTE["muted"])

    photo_card = add_rounded_rect(s, 0.45, 1.0, 7.15, 5.75, PALETTE["card_bg"])
    if os.path.exists(news["image"]):
        photo = s.shapes.add_picture(news["image"], Inches(0.58), Inches(1.13), width=Inches(6.9))
    else:
        photo = add_text(s, "（雪豹照片未找到）", 0.7, 3.0, 6.5, 0.8,
                         size=22, color=PALETTE["muted"], align=PP_ALIGN.CENTER)

    title = add_text(s, news["title"], 7.9, 1.05, 4.9, 1.0,
                     size=34, bold=True, color=PALETTE["ink"])

    fact1_card = add_rounded_rect(s, 7.9, 2.25, 2.25, 0.72, PALETTE["chip_bg"])
    fact1 = add_text(s, news["body1"], 8.0, 2.36, 2.05, 0.5,
                     size=24, bold=True, color=PALETTE["ink"], align=PP_ALIGN.CENTER)

    fact2_card = add_rounded_rect(s, 10.4, 2.25, 2.4, 0.72, PALETTE["chip_bg"])
    fact2 = add_text(s, news["body2"], 10.5, 2.36, 2.2, 0.5,
                     size=24, bold=True, color=PALETTE["ink"], align=PP_ALIGN.CENTER)

    tail_card = add_rounded_rect(s, 7.9, 3.25, 4.9, 0.82, RGBColor(0xFF, 0xF6, 0xE5))
    tail = add_text(s, news["purpose"], 8.05, 3.38, 4.6, 0.55,
                    size=22, bold=True, color=PALETTE["accent"], align=PP_ALIGN.CENTER)

    if os.path.exists(news["animation"]):
        gif = s.shapes.add_picture(news["animation"], Inches(8.35), Inches(4.45), width=Inches(4.0))
        gif_note = add_text(s, "放映时，小雪豹会走起来", 8.35, 6.1, 4.0, 0.35,
                            size=13, color=PALETTE["muted"], align=PP_ALIGN.CENTER)
    else:
        gif = add_text(s, "雪豹走一走", 8.35, 5.1, 4.0, 0.6,
                       size=20, color=PALETTE["muted"], align=PP_ALIGN.CENTER)
        gif_note = gif

    source = add_text(s, "新华社 / 国家林业和草原局 · 2026-10-23", 0.6, 7.0, 6.4, 0.3,
                      size=9, color=PALETTE["muted"])
    page = add_text(s, f"第{slide_num}页，共3页", 9.8, 7.0, 3.0, 0.3,
                    size=10, color=PALETTE["muted"], align=PP_ALIGN.RIGHT)

    add_animation(s, title, delay_ms=200, duration_ms=450)
    add_animation(s, fact1_card, delay_ms=450, duration_ms=400)
    add_animation(s, fact2_card, delay_ms=650, duration_ms=400)
    add_animation(s, tail_card, delay_ms=850, duration_ms=450)
    return s


def build_news_slide(prs, news, slide_num):
    if news.get("special") == "snow_leopard":
        return build_snow_leopard_slide(prs, news, slide_num)

    s = prs.slides.add_slide(prs.slide_layouts[6])

    # Background
    bg = add_rect(s, 0, 0, 13.333, 7.5, PALETTE["bg"])

    # Chip
    chip = add_rounded_rect(s, 0.6, 0.45, 1.1, 0.45, PALETTE["chip_bg"])
    chip_text = add_text(s, news["chip"], 0.6, 0.45, 1.1, 0.45,
             size=14, bold=True, color=PALETTE["ink"],
             align=PP_ALIGN.CENTER)

    # Topic
    topic = add_text(s, news["topic"], 1.9, 0.45, 3.0, 0.45,
             size=14, color=PALETTE["muted"])

    # Image
    img_path = news["image"]
    if os.path.exists(img_path):
        img_card = add_rounded_rect(s, 0.5, 1.1, 6.7, 5.7, PALETTE["card_bg"])
        img_shape = s.shapes.add_picture(img_path, Inches(0.6), Inches(1.2),
                                          Inches(6.5), Inches(5.5))
    else:
        img_card = add_rounded_rect(s, 0.6, 1.2, 6.5, 5.5, PALETTE["chip_bg"])
        img_shape = add_text(s, "（图片未找到）", 0.6, 3.0, 6.5, 1.0,
                 size=20, color=PALETTE["muted"], align=PP_ALIGN.CENTER)

    # Title
    title = add_text(s, news["title"], 7.5, 1.2, 5.3, 1.2,
             size=36, bold=True, color=PALETTE["ink"])

    # Body 1
    body1 = add_text(s, news["body1"], 7.5, 2.6, 5.3, 0.8,
             size=22, color=PALETTE["ink"])

    # Body 2
    body2 = add_text(s, news["body2"], 7.5, 3.5, 5.3, 0.8,
             size=22, color=PALETTE["ink"])

    # Purpose
    purpose_card = add_rounded_rect(s, 7.5, 4.6, 5.3, 0.75, RGBColor(0xFF, 0xF6, 0xE5))
    purpose = add_text(s, news["purpose"], 7.6, 4.65, 5.1, 0.65,
             size=18, bold=True, color=PALETTE["accent"])

    # Footer
    footer_left = add_text(s, "小小新闻播报 · 2026年10月27日 星期二", 0.6, 7.0, 6.0, 0.35,
             size=10, color=PALETTE["muted"])
    footer_right = add_text(s, f"第{slide_num}页，共3页", 7.5, 7.0, 5.3, 0.35,
             size=10, color=PALETTE["muted"], align=PP_ALIGN.RIGHT)

    # Add animations with staggered delays
    add_animation(s, bg, delay_ms=0, duration_ms=300)
    add_animation(s, chip, delay_ms=200, duration_ms=400)
    add_animation(s, topic, delay_ms=300, duration_ms=400)
    add_animation(s, img_card, delay_ms=400, duration_ms=500)
    add_animation(s, title, delay_ms=600, duration_ms=500)
    add_animation(s, body1, delay_ms=800, duration_ms=500)
    add_animation(s, body2, delay_ms=1000, duration_ms=500)
    add_animation(s, purpose_card, delay_ms=1200, duration_ms=500)
    add_animation(s, footer_left, delay_ms=1400, duration_ms=400)
    add_animation(s, footer_right, delay_ms=1400, duration_ms=400)


def save_pptx(prs, path):
    prs.save(path)
    import zipfile as _zf
    import re as _re
    tmp = path + ".tmp"
    try:
        with _zf.ZipFile(path, "r") as zin, _zf.ZipFile(tmp, "w") as zout:
            for item in zin.infolist():
                if item.filename.lower().startswith("docprops/thumbnail"):
                    continue
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    content = data.decode("utf-8")
                    content = _re.sub(
                        r'<Override[^>]*PartName="/docProps/thumbnail[^"]*"[^>]*/?>',
                        '', content, flags=_re.IGNORECASE)
                    data = content.encode("utf-8")
                if item.filename == "_rels/.rels":
                    content = data.decode("utf-8")
                    content = _re.sub(
                        r'<Relationship[^>]*Target="docProps/thumbnail[^"]*"[^>]*/?>',
                        '', content, flags=_re.IGNORECASE)
                    data = content.encode("utf-8")
                zout.writestr(item, data)
        os.replace(tmp, path)
    except Exception as e:
        print("warning: failed to strip template thumbnail:", e)
        try:
            if os.path.exists(tmp):
                os.remove(tmp)
        except Exception:
            pass


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for i, news in enumerate(NEWS, 1):
        build_news_slide(prs, news, i)

    out_dir = os.path.join(SCRIPT_DIR, "daily-news", "2026-10-27")
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, "今日新闻-2026-10-27.pptx")
    save_pptx(prs, out)
    print("wrote", out)


if __name__ == "__main__":
    main()

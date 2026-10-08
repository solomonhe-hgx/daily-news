"""Generate 今日新闻-2026-10-08.pptx — 1 slide with 3 news cards."""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- Palette ----------
PALETTE = {
    "dominant":  RGBColor(0x26, 0x38, 0x4A),
    "green_bg":  RGBColor(0xEA, 0xF8, 0xEF),
    "blue_bg":   RGBColor(0xEA, 0xF5, 0xFF),
    "peach_bg":  RGBColor(0xFF, 0xF0, 0xE7),
    "yellow":    RGBColor(0xFF, 0xD7, 0x5E),
    "red":       RGBColor(0xD9, 0x4F, 0x62),
    "ink":       RGBColor(0x26, 0x38, 0x4A),
    "muted":     RGBColor(0x65, 0x74, 0x82),
    "white":     RGBColor(0xFF, 0xFF, 0xFF),
    "cream":     RGBColor(0xFF, 0xF9, 0xE8),
}

FONT_HEAD = "PingFang SC"
FONT_BODY = "PingFang SC"
FONT_HEAD_EA = "PingFang SC"
FONT_BODY_EA = "PingFang SC"

BASE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(BASE, "assets")


def set_fonts(run, latin, ea):
    run.font.name = latin
    rPr = run._r.get_or_add_rPr()
    ea_el = rPr.find(qn("a:ea"))
    if ea_el is None:
        ea_el = rPr.makeelement(qn("a:ea"), {"typeface": ea})
        rPr.append(ea_el)
    else:
        ea_el.set("typeface", ea)


def add_text(slide, text, x, y, w, h, *, size=14, bold=False,
             color=None, align=PP_ALIGN.LEFT, font=None, font_ea=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font or FONT_BODY
    r.font.size = Pt(size)
    r.font.bold = bold
    if color is not None:
        r.font.color.rgb = color
    set_fonts(r, font or FONT_BODY, font_ea or FONT_BODY_EA)
    return tb


def add_rect(slide, x, y, w, h, fill):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_pill(slide, x, y, w, h, fill):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    # Make it very round
    shp.adjustments[0] = 0.5
    return shp


def save_pptx(prs, path):
    prs.save(path)
    import zipfile as _zf
    import re as _re
    tmp = path + ".tmp"
    try:
        with _zf.ZipFile(path, "r") as zin, _zf.ZipFile(tmp, "w") as zout:
            for item in zin.infolist():
                if item.filename.lower().startswith("docprops/thumbnail"):
                    continue
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    content = data.decode("utf-8")
                    content = _re.sub(
                        r'<Override[^>]*PartName="/docProps/thumbnail[^"]*"[^>]*/?>',
                        '', content, flags=_re.IGNORECASE)
                    data = content.encode("utf-8")
                if item.filename == "_rels/.rels":
                    content = data.decode("utf-8")
                    content = _re.sub(
                        r'<Relationship[^>]*Target="docProps/thumbnail[^"]*"[^>]*/?>',
                        '', content, flags=_re.IGNORECASE)
                    data = content.encode("utf-8")
                zout.writestr(item, data)
        os.replace(tmp, path)
    except Exception as e:
        print("warning: failed to strip thumbnail:", e)
        try:
            if os.path.exists(tmp):
                os.remove(tmp)
        except Exception:
            pass


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ---- Slide 1: Three news cards ----
    s = prs.slides.add_slide(prs.slide_layouts[6])

    # Background
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                            Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = PALETTE["cream"]
    bg.line.fill.background()

    # Title
    add_text(s, "小小新闻", 0.5, 0.25, 4.0, 0.8,
             size=40, bold=True, color=PALETTE["dominant"],
             font=FONT_HEAD, font_ea=FONT_HEAD_EA)

    # Weather pill (top-right)
    add_pill(s, 8.2, 0.3, 4.8, 0.55, PALETTE["white"])
    add_text(s, "10月8日 星期四 上海 晴间多云 17—24℃", 8.4, 0.33, 4.4, 0.5,
             size=16, bold=True, color=PALETTE["dominant"],
             align=PP_ALIGN.CENTER, font=FONT_HEAD, font_ea=FONT_HEAD_EA)

    # News data
    news = [
        {
            "num": "1",
            "bg": PALETTE["green_bg"],
            "img": os.path.join(ASSETS, "space-national-day.jpg"),
            "title": "神舟二十三号航天员从太空送来国庆祝福",
            "body": "10月1日国庆节，三名航天员叔叔从天宫空间站给大家送来祝福。他们已经在太空住了五个多月，一边做实验一边给大家拍视频！",
            "ask": "你想到太空里去住几天？",
            "source": "央视网 · 10月1日",
            "img_note": "航天员在太空站里向大家挥手送祝福",
        },
        {
            "num": "2",
            "bg": PALETTE["blue_bg"],
            "img": os.path.join(ASSETS, "panda-atlanta.jpg"),
            "title": '大熊猫\u201c平平\u201d\u201c福双\u201d坐飞机去了美国',
            "body": '9月27日，两只大熊猫坐专机飞到了美国亚特兰大。它们要在那里住十年，每天吃竹子、睡觉、爬树，过几天就会和大家见面啦！',
            "ask": "你想给它们取什么英文名？",
            "source": "新华社 · 9月28日",
            "img_note": "两只大熊猫正坐在地上开心地吃竹子",
        },
        {
            "num": "3",
            "bg": PALETTE["peach_bg"],
            "img": os.path.join(ASSETS, "fast-telescope.jpg"),
            "title": '贵州大山里的\u201c天眼\u201d找到了上千颗脉冲星',
            "body": '贵州的大山里，有一个全世界最大的\u201c大耳朵\u201d。它叫中国天眼，已经工作十年了，找到了一千多颗脉冲星\u2014\u2014它们是宇宙里会\u201c发信号\u201d的小星星！',
            "ask": "你听到过星星的声音吗？",
            "source": "新华社 · 9月25日",
            "img_note": '贵州大山里有一个巨大的银色\u201c碗\u201d，就是中国天眼',
        },
    ]

    # Layout: 3 cards
    card_w = 3.85
    gap = 0.25
    start_x = 0.5
    card_y = 1.15
    card_h = 6.0

    for i, n in enumerate(news):
        x = start_x + i * (card_w + gap)

        # Card background
        add_rect(s, x, card_y, card_w, card_h, n["bg"])

        # Number chip
        chip = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(x + 0.15), Inches(card_y + 0.15),
                                  Inches(0.42), Inches(0.42))
        chip.fill.solid()
        chip.fill.fore_color.rgb = PALETTE["yellow"]
        chip.line.fill.background()
        chip.adjustments[0] = 0.25
        add_text(s, n["num"], x + 0.15, card_y + 0.15, 0.42, 0.42,
                 size=18, bold=True, color=PALETTE["dominant"],
                 align=PP_ALIGN.CENTER, font=FONT_HEAD, font_ea=FONT_HEAD_EA)

        # Image
        img_y = card_y + 0.7
        img_h = 2.2
        pic = s.shapes.add_picture(n["img"],
                                   Inches(x + 0.15), Inches(img_y),
                                   width=Inches(card_w - 0.3))
        # Clip to rounded rect
        sppr = pic._element.spPr
        prstGeom = sppr.find(qn("a:prstGeom"))
        if prstGeom is not None:
            prstGeom.set("prst", "roundRect")

        # Photo note (below image)
        note_y = img_y + img_h + 0.05
        add_text(s, "图里：" + n["img_note"], x + 0.2, note_y, card_w - 0.4, 0.35,
                 size=10, bold=True, color=PALETTE["muted"],
                 align=PP_ALIGN.CENTER, font=FONT_BODY, font_ea=FONT_BODY_EA)

        # Title
        title_y = note_y + 0.38
        add_text(s, n["title"], x + 0.2, title_y, card_w - 0.4, 0.7,
                 size=16, bold=True, color=PALETTE["dominant"],
                 font=FONT_HEAD, font_ea=FONT_HEAD_EA)

        # Body
        body_y = title_y + 0.65
        add_text(s, n["body"], x + 0.2, body_y, card_w - 0.4, 1.2,
                 size=12, color=PALETTE["ink"],
                 font=FONT_BODY, font_ea=FONT_BODY_EA)

        # Ask question
        ask_y = body_y + 1.15
        add_text(s, n["ask"], x + 0.2, ask_y, card_w - 0.4, 0.35,
                 size=13, bold=True, color=PALETTE["red"],
                 font=FONT_BODY, font_ea=FONT_BODY_EA)

        # Source
        src_y = ask_y + 0.35
        add_text(s, n["source"], x + 0.2, src_y, card_w - 0.4, 0.3,
                 size=9, color=PALETTE["muted"],
                 font=FONT_BODY, font_ea=FONT_BODY_EA)

    # Footer
    add_text(s, "小主播：选一条你最喜欢的新闻讲给大家听。",
             0.5, 7.15, 8.0, 0.3,
             size=11, bold=True, color=PALETTE["muted"],
             font=FONT_BODY, font_ea=FONT_BODY_EA)

    out_path = os.path.join(BASE, "今日新闻-2026-10-08.pptx")
    save_pptx(prs, out_path)
    print("wrote", out_path)


if __name__ == "__main__":
    main()

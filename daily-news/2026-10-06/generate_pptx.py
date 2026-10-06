"""
generate_pptx.py — Generate 今日新闻-2026-10-06.pptx
Creates a 2-slide PPTX with 3 child-friendly news items.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

# ---------- Palette ----------
PALETTE = {
    "cream":    RGBColor(0xFF, 0xF9, 0xE8),
    "green":    RGBColor(0xEA, 0xF8, 0xEF),
    "blue":     RGBColor(0xEA, 0xF5, 0xFF),
    "peach":    RGBColor(0xFF, 0xF0, 0xE7),
    "ink":      RGBColor(0x26, 0x38, 0x4A),
    "red":      RGBColor(0xFF, 0x64, 0x75),
    "yellow":   RGBColor(0xFF, 0xD7, 0x5E),
    "white":    RGBColor(0xFF, 0xFF, 0xFF),
    "muted":    RGBColor(0x65, 0x74, 0x82),
}

FONT_HEAD = "PingFang SC"
FONT_BODY = "PingFang SC"
FONT_HEAD_EA = "PingFang SC"
FONT_BODY_EA = "PingFang SC"

ASSETS = os.path.join(os.path.dirname(__file__), "assets")


def add_text(slide, text, x, y, w, h, *,
             size=14, bold=False, color=None, align=PP_ALIGN.LEFT,
             font=None, font_ea=None, line_spacing=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing:
        p.line_spacing = Pt(line_spacing)
    r = p.add_run()
    r.text = text
    r.font.name = font or FONT_BODY
    r.font.size = Pt(size)
    r.font.bold = bold
    if color is not None:
        r.font.color.rgb = color
    # Set East Asian font
    from pptx.oxml.ns import qn
    rPr = r._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {"typeface": font_ea or FONT_BODY_EA})
        rPr.append(ea)
    else:
        ea.set("typeface", font_ea or FONT_BODY_EA)
    return tb


def add_multiline_text(slide, lines, x, y, w, h, *,
                       size=14, bold=False, color=None, align=PP_ALIGN.LEFT,
                       font=None, font_ea=None, line_spacing=None):
    """Add text with multiple paragraphs/lines."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)

    for i, line_text in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = Pt(line_spacing)
        p.space_after = Pt(2)
        r = p.add_run()
        r.text = line_text
        r.font.name = font or FONT_BODY
        r.font.size = Pt(size)
        r.font.bold = bold
        if color is not None:
            r.font.color.rgb = color
        from pptx.oxml.ns import qn
        rPr = r._r.get_or_add_rPr()
        ea = rPr.find(qn("a:ea"))
        if ea is None:
            ea = rPr.makeelement(qn("a:ea"), {"typeface": font_ea or FONT_BODY_EA})
            rPr.append(ea)
        else:
            ea.set("typeface", font_ea or FONT_BODY_EA)
    return tb


def add_rect(slide, x, y, w, h, fill, line=None, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
    shp.shadow.inherit = False
    return shp


def add_image(slide, img_path, x, y, w, h):
    shp = slide.shapes.add_picture(img_path, Inches(x), Inches(y), Inches(w), Inches(h))
    return shp


def add_news_card(slide, num, img_path, title, body_lines, ask_text, source_text,
                  x, y, w, h, bg_color):
    """Add a news card with image, title, body, and source."""
    # Card background
    add_rect(slide, x, y, w, h, bg_color, radius=True)

    # Number badge
    badge_size = 0.4
    add_rect(slide, x + 0.15, y + 0.15, badge_size, badge_size,
             PALETTE["yellow"], radius=True)
    add_text(slide, str(num), x + 0.15, y + 0.15, badge_size, badge_size,
             size=16, bold=True, color=PALETTE["ink"], align=PP_ALIGN.CENTER)

    # Image
    img_h = 1.8
    img_margin = 0.15
    add_image(slide, img_path, x + img_margin, y + 0.65, w - 2 * img_margin, img_h)

    # Photo note overlay
    photo_note_y = y + 0.65 + img_h - 0.35
    add_rect(slide, x + img_margin + 0.1, photo_note_y,
             w - 2 * img_margin - 0.2, 0.28,
             RGBColor(0x26, 0x38, 0x4A), radius=True)

    # Title
    title_y = y + 0.65 + img_h + 0.05
    add_text(slide, title, x + 0.2, title_y, w - 0.4, 0.5,
             size=14, bold=True, color=PALETTE["ink"], line_spacing=18)

    # Body
    body_y = title_y + 0.5
    add_multiline_text(slide, body_lines, x + 0.2, body_y, w - 0.4, 0.8,
                       size=11, color=PALETTE["ink"], line_spacing=15)

    # Ask text
    ask_y = body_y + 0.7
    add_text(slide, ask_text, x + 0.2, ask_y, w - 0.4, 0.3,
             size=11, bold=True, color=PALETTE["red"], line_spacing=14)

    # Source
    source_y = ask_y + 0.25
    add_text(slide, source_text, x + 0.2, source_y, w - 0.4, 0.2,
             size=7, color=PALETTE["muted"], line_spacing=10)


def save_pptx(prs, path):
    """Save PPTX and strip blank placeholder thumbnail."""
    prs.save(path)
    import zipfile as _zf
    import re as _re
    tmp = path + ".tmp"
    try:
        with _zf.ZipFile(path, "r") as zin, _zf.ZipFile(tmp, "w") as zout:
            for item in zin.infolist():
                if item.filename.lower().startswith("docprops/thumbnail"):
                    continue
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    content = data.decode("utf-8")
                    content = _re.sub(
                        r'<Override[^>]*PartName="/docProps/thumbnail[^"]*"[^>]*/?>',
                        '', content, flags=_re.IGNORECASE
                    )
                    data = content.encode("utf-8")
                if item.filename == "_rels/.rels":
                    content = data.decode("utf-8")
                    content = _re.sub(
                        r'<Relationship[^>]*Target="docProps/thumbnail[^"]*"[^>]*/?>',
                        '', content, flags=_re.IGNORECASE
                    )
                    data = content.encode("utf-8")
                zout.writestr(item, data)
        os.replace(tmp, path)
    except Exception as e:
        print("warning: failed to strip thumbnail:", e)
        try:
            if os.path.exists(tmp):
                os.remove(tmp)
        except Exception:
            pass


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ===== SLIDE 1: News 1 + News 2 =====
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    # Background
    add_rect(s1, 0, 0, 13.333, 7.5, PALETTE["cream"])

    # Title
    add_text(s1, "小小新闻", 0.5, 0.3, 4, 0.8,
             size=36, bold=True, color=PALETTE["ink"])

    # Weather meta (upper right)
    add_rect(s1, 8.5, 0.35, 4.5, 0.5, PALETTE["white"], radius=True)
    add_text(s1, "10月6日 星期二 上海 晴 15—23℃", 8.6, 0.38, 4.3, 0.45,
             size=14, bold=True, color=PALETTE["ink"], align=PP_ALIGN.CENTER)

    # News 1 - Aurora (green card)
    add_news_card(
        s1, num=1,
        img_path=os.path.join(ASSETS, "aurora.jpg"),
        title="黑龙江夜空出现美丽的绿色极光",
        body_lines=[
            "10月4日晚上，黑龙江好几个地方看到了极光。",
            "绿色和紫色的光像彩带一样飘在天上。",
        ],
        ask_text="你觉得极光像什么？",
        source_text="新华社 · 10月5日  图：AI生成",
        x=0.4, y=1.2, w=6.1, h=5.8,
        bg_color=PALETTE["green"]
    )

    # News 2 - Robots (blue card)
    add_news_card(
        s1, num=2,
        img_path=os.path.join(ASSETS, "robots.jpg"),
        title="安徽商场里机器人给小朋友做爆米花",
        body_lines=[
            "国庆假期，安徽的商场里来了机器人店员。",
            "它们会做爆米花、还会和小朋友一起踢球！",
        ],
        ask_text="你想和机器人玩什么？",
        source_text="新华社 · 10月4日  图：AI生成",
        x=6.8, y=1.2, w=6.1, h=5.8,
        bg_color=PALETTE["blue"]
    )

    # ===== SLIDE 2: News 3 + Footer =====
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s2, 0, 0, 13.333, 7.5, PALETTE["cream"])

    # Title
    add_text(s2, "小小新闻", 0.5, 0.3, 4, 0.8,
             size=36, bold=True, color=PALETTE["ink"])

    # Weather meta (upper right)
    add_rect(s2, 8.5, 0.35, 4.5, 0.5, PALETTE["white"], radius=True)
    add_text(s2, "10月6日 星期二 上海 晴 15—23℃", 8.6, 0.38, 4.3, 0.45,
             size=14, bold=True, color=PALETTE["ink"], align=PP_ALIGN.CENTER)

    # News 3 - Space station (peach card, centered)
    add_news_card(
        s2, num=3,
        img_path=os.path.join(ASSETS, "space-station.jpg"),
        title="神舟二十三号航天员在太空住了四个多月",
        body_lines=[
            "三名航天员叔叔在太空站已经住了四个多月。",
            "他们还在太空过了中秋节，一边看月亮一边做实验！",
        ],
        ask_text="你想到太空住多久？",
        source_text="央视网 · 10月5日  图：AI生成",
        x=3.6, y=1.2, w=6.1, h=5.0,
        bg_color=PALETTE["peach"]
    )

    # Footer
    add_text(s2, "小主播：选一条你最喜欢的新闻讲给大家听。",
             0.5, 6.7, 8, 0.4,
             size=12, bold=True, color=PALETTE["muted"])

    # Save
    out_path = os.path.join(os.path.dirname(__file__), "今日新闻-2026-10-06.pptx")
    save_pptx(prs, out_path)
    print(f"wrote {out_path}")


if __name__ == "__main__":
    main()

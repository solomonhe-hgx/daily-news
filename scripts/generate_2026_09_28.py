from pathlib import Path
from shutil import copy2
import os, re, zipfile
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT=Path(__file__).resolve().parents[1]
ASSETS=ROOT/"daily-news/2026-09-28/assets"
OUT=ROOT/"outputs/今日新闻-2026-09-28.pptx"
WEB=ROOT/"daily-news/2026-09-28/今日新闻-2026-09-28.pptx"
FONT="PingFang SC"
C={"ink":RGBColor(0x26,0x38,0x4A),"red":RGBColor(0xFF,0x64,0x75),"yellow":RGBColor(0xFF,0xD7,0x5E),"green":RGBColor(0xEA,0xF8,0xEF),"blue":RGBColor(0xEA,0xF5,0xFF),"peach":RGBColor(0xFF,0xF0,0xE7),"cream":RGBColor(0xFF,0xF9,0xE8),"white":RGBColor(0xFF,0xFF,0xFF),"muted":RGBColor(0x61,0x70,0x80)}
BUDGET={"title":18,"body":38,"tip":22,"small":80}

def guard(value,key):
    if len(value)>BUDGET[key]: raise ValueError(f"text too long for {key}: {value}")
    return value

def set_font(run,size,color,bold=False):
    run.font.name=FONT; run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=color
    rpr=run._r.get_or_add_rPr(); ea=rpr.find(qn("a:ea"))
    if ea is None: ea=rpr.makeelement(qn("a:ea"),{"typeface":FONT}); rpr.append(ea)
    else: ea.set("typeface",FONT)

def text(slide,value,x,y,w,h,size,color=None,bold=False,align=PP_ALIGN.LEFT,budget="body"):
    value=guard(value,budget); box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame
    tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.NONE; tf.margin_left=Inches(.04); tf.margin_right=Inches(.04); tf.margin_top=Inches(.03); tf.margin_bottom=Inches(.03)
    p=tf.paragraphs[0]; p.alignment=align; p.space_after=Pt(0); r=p.add_run(); r.text=value; set_font(r,size,color or C["ink"],bold); return box

def shape(slide,kind,x,y,w,h,fill):
    s=slide.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h)); s.fill.solid(); s.fill.fore_color.rgb=fill; s.line.fill.background(); s.shadow.inherit=False; return s

def rect(slide,x,y,w,h,fill,rounded=True): return shape(slide,MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,x,y,w,h,fill)

def photo(slide,path,x,y,w,h):
    with Image.open(path) as im: iw,ih=im.size
    pic=slide.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h)); fr=w/h; ir=iw/ih
    if ir>fr: c=(1-fr/ir)/2; pic.crop_left=c; pic.crop_right=c
    elif ir<fr: c=(1-ir/fr)/2; pic.crop_top=c; pic.crop_bottom=c
    return pic

def meta(slide):
    rect(slide,8.15,.34,4.5,.52,C["white"]); text(slide,"9月28日 星期一　上海 小雨 23—25℃",8.33,.46,4.15,.23,12,C["ink"],True,PP_ALIGN.CENTER,"small")

def source(slide,value): text(slide,value,.65,7.05,9.8,.2,8,C["muted"],False,budget="small")

def save_pptx(prs,path):
    prs.save(path); tmp=str(path)+".tmp"
    with zipfile.ZipFile(path,"r") as zin, zipfile.ZipFile(tmp,"w") as zout:
        for item in zin.infolist():
            if item.filename.lower().startswith("docprops/thumbnail"): continue
            data=zin.read(item.filename)
            if item.filename=="[Content_Types].xml": data=re.sub(r'<Override[^>]*PartName="/docProps/thumbnail[^"]*"[^>]*/?>','',data.decode(),flags=re.I).encode()
            if item.filename=="_rels/.rels": data=re.sub(r'<Relationship[^>]*Target="docProps/thumbnail[^"]*"[^>]*/?>','',data.decode(),flags=re.I).encode()
            zout.writestr(item,data)
    os.replace(tmp,path)

def add_card(slide,x,bg,image,title,body,question,num):
    rect(slide,x,1.25,5.8,5.45,bg); photo(slide,image,x+.16,1.42,5.48,2.72)
    rect(slide,x+.32,1.58,.55,.55,C["yellow"]); text(slide,str(num),x+.32,1.72,.55,.24,15,C["ink"],True,PP_ALIGN.CENTER,"small")
    text(slide,title,x+.28,4.35,5.2,.5,25,C["ink"],True,budget="title")
    text(slide,body,x+.28,4.93,5.2,.7,19,C["ink"],False,budget="body")
    text(slide,question,x+.28,5.87,5.2,.35,16,C["red"],True,budget="tip")

def build():
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    # Slide 1: two current stories
    s=prs.slides.add_slide(prs.slide_layouts[6]); rect(s,0,0,13.333,7.5,C["cream"],False)
    text(s,"小小新闻",.65,.35,4.3,.55,28,C["ink"],True,budget="title"); meta(s)
    add_card(s,.62,C["green"],ASSETS/"pandas.jpg","熊猫坐飞机","平平和福双到了新家。\n它们先要好好休息。","你想送熊猫什么礼物？",1)
    add_card(s,6.9,C["blue"],ASSETS/"science-letters.jpg","一封科学来信","科学家写信给小朋友：\n多问为什么，勇敢想一想。","你最想问什么？",2)
    source(s,"新闻：新华社、中国科学院 · 2026年9月28日　图片：央视新闻、中国科学院")
    text(s,"1/2",11.75,7.02,.65,.22,9,C["muted"],True,PP_ALIGN.RIGHT,"small")
    # Slide 2: one image-led story
    s=prs.slides.add_slide(prs.slide_layouts[6]); rect(s,0,0,13.333,7.5,C["cream"],False)
    text(s,"小小新闻",.65,.35,4.3,.55,28,C["ink"],True,budget="title"); meta(s)
    photo(s,ASSETS/"tide.jpg",.65,1.28,7.3,5.15); rect(s,.88,1.52,.58,.58,C["yellow"]); text(s,"3",.88,1.67,.58,.24,16,C["ink"],True,PP_ALIGN.CENTER,"small")
    text(s,"潮水像蝴蝶",8.42,1.65,4.1,.78,32,C["ink"],True,budget="title")
    text(s,"三股潮水碰在一起，\n像一只大大的蝴蝶。",8.42,2.8,4.1,1.25,24,C["ink"],False,budget="body")
    rect(s,8.42,4.5,4.15,.82,C["white"]); text(s,"你觉得它还像什么？",8.64,4.72,3.72,.32,18,C["red"],True,budget="tip")
    text(s,"选一条你最喜欢的新闻讲给大家听。",8.42,5.75,4.0,.6,15,C["muted"],True,budget="body")
    source(s,"新闻与图片：中新网 · 2026年9月27日"); text(s,"2/2",11.75,7.02,.65,.22,9,C["muted"],True,PP_ALIGN.RIGHT,"small")
    return prs

def main():
    OUT.parent.mkdir(parents=True,exist_ok=True); WEB.parent.mkdir(parents=True,exist_ok=True)
    save_pptx(build(),OUT); copy2(OUT,WEB); print(OUT); print(WEB)

if __name__=="__main__": main()

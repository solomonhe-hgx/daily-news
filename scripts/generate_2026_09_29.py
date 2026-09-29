from pathlib import Path
from shutil import copy2
import os, re, zipfile
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "daily-news/2026-09-29/assets"
OUT = ROOT / "outputs/今日新闻-2026-09-29.pptx"
WEB = ROOT / "daily-news/2026-09-29/今日新闻-2026-09-29.pptx"

FONT = "PingFang SC"
C = {
    "ink": RGBColor(0x26, 0x38, 0x4A),
    "red": RGBColor(0xFF, 0x64, 0x75),
    "yellow": RGBColor(0xFF, 0xD7, 0x5E),
    "green": RGBColor(0xEA, 0xF8, 0xEF),
    "blue": RGBColor(0xEA, 0xF5, 0xFF),
    "peach": RGBColor(0xFF, 0xF0, 0xE7),
    "cream": RGBColor(0xFF, 0xF9, 0xE8),
    "white": RGBColor(0xFF, 0xFF, 0xFF),
    "muted": RGBColor(0x61, 0x70, 0x80),
}
BUDGET = {"title": 18, "body": 60, "tip": 22, "small": 90}


def guard(value, key):
    if len(value) > BUDGET[key]:
        raise ValueError(f"text too long for {key}: {value}")
    return value


def set_font(run, size, color, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = rpr.makeelement(qn("a:ea"), {"typeface": FONT})
        rpr.append(ea)
    else:
        ea.set("typeface", FONT)


def text(slide, value, x, y, w, h, size, color=None, bold=False, align=PP_ALIGN.LEFT, budget="body"):
    value = guard(value, budget)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    r = p.add_run()
    r.text = value
    set_font(r, size, color or C["ink"], bold)
    return box


def shape(slide, kind, x, y, w, h, fill):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def rect(slide, x, y, w, h, fill, rounded=True):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h, fill)


def photo(slide, path, x, y, w, h):
    with Image.open(path) as im:
        iw, ih = im.size
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    fr = w / h
    ir = iw / ih
    if ir > fr:
        c = (1 - fr / ir) / 2
        pic.crop_left = c
        pic.crop_right = c
    elif ir < fr:
        c = (1 - ir / fr) / 2
        pic.crop_top = c
        pic.crop_bottom = c
    return pic


def meta(slide):
    rect(slide, 7.35, 0.34, 5.3, 0.52, C["white"])
    text(slide, "9月29日 星期二　上海 小雨 22—26℃", 7.53, 0.46, 4.95, 0.23, 11, C["ink"], True, PP_ALIGN.CENTER, "small")


def source(slide, value):
    text(slide, value, 0.65, 7.05, 12.0, 0.2, 8, C["muted"], False, budget="small")


def save_pptx(prs, path):
    prs.save(path)
    tmp = str(path) + ".tmp"
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w") as zout:
        for item in zin.infolist():
            if item.filename.lower().startswith("docprops/thumbnail"):
                continue
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = re.sub(
                    r'<Override[^>]*PartName="/docProps/thumbnail[^"]*"[^>]*/?>',
                    "",
                    data.decode(),
                    flags=re.I,
                ).encode()
            if item.filename == "_rels/.rels":
                data = re.sub(
                    r'<Relationship[^>]*Target="docProps/thumbnail[^"]*"[^>]*/?>',
                    "",
                    data.decode(),
                    flags=re.I,
                ).encode()
            zout.writestr(item, data)
    os.replace(tmp, path)


def add_card(slide, x, bg, image, image_note, title, body, question, num):
    rect(slide, x, 1.25, 4.0, 5.45, bg)
    photo(slide, image, x + 0.16, 1.42, 3.68, 2.42)
    rect(slide, x + 0.32, 1.58, 0.55, 0.55, C["yellow"])
    text(slide, str(num), x + 0.32, 1.72, 0.55, 0.24, 15, C["ink"], True, PP_ALIGN.CENTER, "small")
    rect(slide, x + 0.16, 3.32, 3.68, 0.4, C["ink"])
    text(slide, image_note, x + 0.28, 3.4, 3.45, 0.22, 10, C["white"], True, PP_ALIGN.CENTER, "tip")
    text(slide, title, x + 0.18, 3.95, 3.65, 0.85, 22, C["ink"], True, budget="title")
    text(slide, body, x + 0.18, 4.83, 3.65, 1.4, 16, C["ink"], False, budget="body")
    rect(slide, x + 0.18, 6.0, 3.65, 0.5, C["white"])
    text(slide, question, x + 0.28, 6.12, 3.45, 0.28, 15, C["red"], True, budget="tip")


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 13.333, 7.5, C["cream"], False)
    text(s, "小小新闻", 0.65, 0.35, 4.3, 0.55, 28, C["ink"], True, budget="title")
    meta(s)

    add_card(
        s, 0.55, C["green"], ASSETS / "pandas.jpg",
        "图里：熊猫正坐大飞机去新家",
        "熊猫平平福双坐飞机去美国",
        "它们从成都出发，\n飞到美国的新动物园。\n先好好休息，再见面。",
        "你想给它们说什么？",
        1,
    )
    add_card(
        s, 4.66, C["blue"], ASSETS / "tide-qiantang.jpg",
        "图里：钱塘江一条白浪推着走",
        "钱塘江大潮像白线一样",
        "一条白白浪花，\n推着江面往前走。\n小朋友们大声说：好厉害！",
        "你觉得它像什么？",
        2,
    )
    add_card(
        s, 8.77, C["peach"], ASSETS / "moon-autumn.jpg",
        "图里：一轮又大又圆的月亮挂在天上",
        "中秋的大月亮又圆又大",
        "今年最圆的月亮，\n比中秋晚两天才来。\n抬头看一看，有小白兔。",
        "你看到了什么？",
        3,
    )
    source(s, "新闻：新华社、央视新闻 · 2026年9月24—28日　图片：央视新闻、视觉中国、中评网")
    text(s, "1/1", 11.75, 7.02, 0.65, 0.22, 9, C["muted"], True, PP_ALIGN.RIGHT, "small")
    return prs


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    WEB.parent.mkdir(parents=True, exist_ok=True)
    save_pptx(build(), OUT)
    copy2(OUT, WEB)
    print(OUT)
    print(WEB)


if __name__ == "__main__":
    main()
